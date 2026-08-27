from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as RLTable,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DOCX_PATH = OUT / "Coter-Pro_Dossier-Clinicas.docx"
PDF_PATH = ROOT / "output" / "pdf" / "Coter-Pro_Dossier-Clinicas.pdf"
BRAND_IMAGE = ROOT / "www" / "og-image.png"


# narrative_proposal preset, with a named Coter brand override for headings/callouts.
NAVY = "17234D"
INK = "17234D"
PURPLE = "7147F5"
CYAN = "2CCFCA"
MUTED = "5D6885"
PALE_BLUE = "EEF2FF"
PALE_PURPLE = "F3F0FF"
PALE_CYAN = "E9FAF8"
LIGHT_LINE = "D9E0F0"
WHITE = "FFFFFF"

USABLE_DXA = 9360


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_LINE, size="6", **edges):
    """Apply restrained single borders; edge values are single or nil."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        value = edges.get(edge, "single")
        tag = qn(f"w:{edge}")
        el = tc_borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tc_borders.append(el)
        el.set(qn("w:val"), value)
        if value != "nil":
            el.set(qn("w:sz"), size)
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=120):
    """Make table DXA geometry deterministic: tblW, tblGrid and every tcW agree."""
    assert sum(widths) == USABLE_DXA, (sum(widths), widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MUTED)


def add_bottom_border(paragraph, color=LIGHT_LINE, size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_text(doc, text, style="Normal", align=None, color=None, size=None, bold=None, italic=None, after=None, before=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    if any(v is not None for v in (color, size, bold, italic)):
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_kicker(doc, text, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text.upper())
    set_run_font(run, size=9.5, color=PURPLE, bold=True)
    return p


def add_lead(doc, text, after=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=13.2, color=NAVY)
    return p


def add_section_title(doc, title, intro=None):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(title)
    set_run_font(r, size=16, color=NAVY, bold=True)
    if intro:
        add_lead(doc, intro, after=15)
    return p


def add_feature(doc, title, body, accent=PURPLE):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=13, color=accent, bold=True)
    body_p = doc.add_paragraph()
    body_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body_p.paragraph_format.space_before = Pt(0)
    body_p.paragraph_format.space_after = Pt(8)
    body_p.paragraph_format.line_spacing = 1.22
    body_r = body_p.add_run(body)
    set_run_font(body_r, size=10.8, color=INK)


def add_callout(doc, title, body, fill=PALE_PURPLE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, color=fill)
    set_cell_margins(cell, top=150, start=220, bottom=150, end=220)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    set_run_font(r, size=11.4, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def fill_table_cell(cell, title, body, header=False):
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3 if not header else 0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    set_run_font(r, size=9.4 if header else 10.3, color=WHITE if header else NAVY, bold=True)
    if body:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.12
        r2 = p2.add_run(body)
        set_run_font(r2, size=9.4, color=INK)


def add_data_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell, color=NAVY)
        set_cell_margins(cell, top=120, start=130, bottom=120, end=130)
        fill_table_cell(cell, header, "", header=True)
    mark_header_row(table.rows[0])
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, pair in zip(cells, row):
            if isinstance(pair, tuple):
                title, body = pair
            else:
                title, body = pair, ""
            set_cell_shading(cell, WHITE if index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell, color=LIGHT_LINE)
            set_cell_margins(cell, top=130, start=130, bottom=130, end=130)
            fill_table_cell(cell, title, body)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def page_break(doc):
    doc.add_page_break()


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(29)
    title.font.color.rgb = rgb(NAVY)
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.line_spacing = 1.2

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, PURPLE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = rgb(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.0

    # Header/footer: named Coter brand override consistent across the dossier.
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("COTER PRO  /  DOSSIER PARA CLINICAS")
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    add_bottom_border(p, color=LIGHT_LINE, size="6")

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(footer_table, [7020, 2340])
    left, right = footer_table.rows[0].cells
    for cell in (left, right):
        set_cell_borders(cell, top="nil", left="nil", bottom="nil", right="nil")
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    p_left = left.paragraphs[0]
    p_left.paragraph_format.space_after = Pt(0)
    r_left = p_left.add_run("Coter Pro - Continuidad terapeutica entre sesiones")
    set_run_font(r_left, size=8.5, color=MUTED)
    p_right = right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_after = Pt(0)
    r_right = p_right.add_run("Pagina ")
    set_run_font(r_right, size=8.5, color=MUTED)
    add_page_number(p_right)


def add_cover(doc):
    add_kicker(doc, "Dossier de producto para clinicas", after=12)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Continuidad terapeutica entre sesiones")
    set_run_font(r, size=29, color=NAVY, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = subtitle.add_run(
        "Una plataforma para que el equipo clinico acompañe el proceso entre consultas "
        "con mas estructura, contexto y claridad."
    )
    set_run_font(r2, size=14, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    inline = run.add_picture(str(BRAND_IMAGE), width=Inches(6.5))
    inline._inline.docPr.set("descr", "Imagen de producto de Coter Pro con sus principales capacidades")
    inline._inline.docPr.set("title", "Coter Pro")

    add_callout(
        doc,
        "Una capa de continuidad para la practica clinica.",
        "Coter Pro centraliza check-ins, tareas terapeuticas, escalas, notas y comunicacion "
        "en un entorno pensado para que cada sesion empiece con la historia reciente del paciente.",
        fill=PALE_CYAN,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Preparado para conversaciones de demostracion y piloto  |  Julio de 2026")
    set_run_font(r, size=9.5, color=MUTED, italic=True)


def add_problem_page(doc):
    add_kicker(doc, "Por que existe Coter Pro")
    add_section_title(
        doc,
        "Lo importante no ocurre solo durante la sesion.",
        "Entre consultas aparecen avances, bloqueos, dudas y momentos relevantes. Sin un espacio comun, "
        "esa informacion llega tarde, se reparte en distintos canales o se pierde.",
    )
    add_feature(
        doc,
        "Una historia reciente, no solo el recuerdo de la ultima cita.",
        "Los check-ins breves ayudan a recoger estado de animo, ansiedad, energia y reflexiones cuando son relevantes; "
        "el profesional los revisa dentro del contexto del proceso.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "Tareas que dejan de depender de mensajes y memoria.",
        "El equipo puede proponer una tarea concreta, con una estructura clara para el paciente y una respuesta que se "
        "recupera antes de la siguiente consulta.",
        accent=CYAN,
    )
    add_feature(
        doc,
        "Un canal profesional separado de la mensajeria personal.",
        "Coter concentra la comunicacion vinculada al seguimiento en el mismo lugar que los objetivos, las notas y la actividad reciente.",
        accent=PURPLE,
    )
    add_callout(
        doc,
        "La conversacion clinica sigue siendo el centro.",
        "Coter no sustituye la relacion terapeutica ni el juicio profesional: da una estructura ligera para que el tiempo entre sesiones aporte contexto util.",
        fill=PALE_PURPLE,
    )


def add_workflow_page(doc):
    add_kicker(doc, "Como se utiliza")
    add_section_title(
        doc,
        "Un flujo directo para profesionales y pacientes.",
        "La experiencia esta diseñada para reducir friccion: el profesional define el seguimiento y el paciente accede con una interfaz sencilla.",
    )
    add_data_table(
        doc,
        ["Momento", "Que ocurre", "Resultado para la consulta"],
        [
            (
                "Preparacion",
                "El profesional crea el acceso, define objetivos y asigna tareas o cuestionarios acordes al caso.",
                "Un plan entre sesiones claro y personalizado.",
            ),
            (
                "Entre sesiones",
                "El paciente completa check-ins, ejercicios estructurados y escalas desde su espacio personal.",
                "Informacion ordenada, no mensajes dispersos.",
            ),
            (
                "Antes de la cita",
                "El profesional revisa actividad, respuestas, progreso y alertas de seguimiento desde el panel.",
                "La sesion puede empezar con contexto reciente.",
            ),
        ],
        [1680, 4260, 3420],
    )
    add_feature(
        doc,
        "Pensado para una adopcion gradual.",
        "Una clinica puede empezar con un grupo acotado de profesionales y pacientes, acordar que informacion es util y ajustar el flujo antes de ampliarlo.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "Una experiencia paciente que no abruma.",
        "La persona encuentra una tarea visible, un check-in breve y el historial de su propio progreso. El objetivo es facilitar el uso real, no llenar su dia de pantallas.",
        accent=CYAN,
    )


def add_capabilities_page(doc):
    add_kicker(doc, "Capacidades de producto")
    add_section_title(
        doc,
        "Las herramientas que sostienen el seguimiento.",
        "Coter Pro combina funciones practicas que el equipo puede utilizar segun su modelo terapeutico y sus protocolos internos.",
    )
    add_feature(
        doc,
        "Check-ins emocionales",
        "Registro breve de animo, ansiedad, energia y pensamientos recientes para observar el intervalo entre consultas y enriquecer la conversacion clinica.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "Tareas terapeuticas estructuradas",
        "Biblioteca y tareas personalizadas para trabajo cognitivo-conductual, registros de pensamiento, activacion conductual, exposicion gradual y ejercicios clasicos.",
        accent=CYAN,
    )
    add_feature(
        doc,
        "Escalas y evolucion",
        "Cuestionarios como PHQ-9, GAD-7 y BDI-II con puntuacion e historial temporal, utilizados como apoyo a la valoracion profesional.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "Panel clinico con contexto",
        "Notas, objetivos, mensajes, actividad reciente y alertas de seguimiento reunidos en una unica vista para preparar el encuentro terapeutico.",
        accent=CYAN,
    )
    add_callout(
        doc,
        "Informacion para revisar, no decisiones automaticas.",
        "Las escalas, tendencias y alertas son apoyos de seguimiento. La interpretacion, la priorizacion y cualquier decision clinica pertenecen siempre al profesional responsable.",
        fill=PALE_CYAN,
    )


def add_pilot_page(doc):
    add_kicker(doc, "Piloto con una clinica")
    add_section_title(
        doc,
        "Un despliegue pequeño, medible y adaptable.",
        "La mejor manera de valorar Coter es probarlo con un reto concreto: por ejemplo, mejorar la adherencia a tareas o llegar a cada sesion con una vision mas actualizada.",
    )
    add_data_table(
        doc,
        ["Etapa", "Trabajo conjunto", "Evidencia que se revisa"],
        [
            (
                "Definicion",
                "Seleccionar perfiles, protocolo de uso, responsables y limites del piloto.",
                "Objetivo operativo y criterios de exito acordados.",
            ),
            (
                "Configuracion",
                "Preparar tareas, escalas, mensajes de bienvenida y una orientacion breve al equipo.",
                "Flujo listo para un grupo inicial de pacientes.",
            ),
            (
                "Seguimiento",
                "Revisar uso, fricciones y aprendizaje con una cadencia acordada.",
                "Datos de adopcion y feedback cualitativo del equipo.",
            ),
            (
                "Decision",
                "Valorar ajustes, continuidad o ampliacion del alcance.",
                "Conclusiones utiles para la clinica y para Coter.",
            ),
        ],
        [1560, 4680, 3120],
    )
    add_feature(
        doc,
        "Indicadores que una clinica puede acordar desde el inicio.",
        "Uso sostenido por pacientes, porcentaje de tareas completadas, frecuencia de check-ins, tiempo de preparacion de sesiones y valoracion del equipo sobre la calidad del contexto disponible.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "El piloto no exige cambiar todo de golpe.",
        "El alcance, el numero de profesionales, los perfiles de pacientes y la frecuencia de revision se definen de forma conjunta. Primero se valida la utilidad; despues se decide si tiene sentido escalar.",
        accent=CYAN,
    )


def add_trust_page(doc):
    add_kicker(doc, "Confianza y siguiente paso")
    add_section_title(
        doc,
        "Tecnologia al servicio de una practica responsable.",
        "Coter se ha construido para trabajar con informacion sensible y para preservar la responsabilidad clinica en manos de las personas.",
    )
    add_feature(
        doc,
        "Proteccion de datos sensibles",
        "El producto incorpora cifrado de datos sensibles en base de datos, sesiones seguras, controles de acceso separados entre profesional y paciente y registro de actividad relevante.",
        accent=PURPLE,
    )
    add_feature(
        doc,
        "Uso responsable",
        "Coter no diagnostica, no recomienda tratamientos y no sustituye los protocolos de crisis de la clinica. No debe utilizarse como canal de atencion urgente; cada despliegue debe definir sus circuitos asistenciales y de proteccion de datos.",
        accent=CYAN,
    )
    add_feature(
        doc,
        "Integracion con criterio clinico",
        "Las tareas, escalas, mensajes y alertas se configuran alrededor del enfoque y los limites de cada equipo. La herramienta acompaña el proceso; no impone un modelo terapeutico unico.",
        accent=PURPLE,
    )
    add_callout(
        doc,
        "El siguiente paso: una conversacion de descubrimiento.",
        "Identifiquemos un reto de seguimiento que merezca la pena resolver, el grupo inicial de pacientes y las condiciones necesarias para un piloto seguro y util. Solicita una demostracion en coter.app.",
        fill=PALE_PURPLE,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("COTER PRO  ·  CONTINUIDAD TERAPEUTICA ENTRE SESIONES")
    set_run_font(r, size=10.5, color=NAVY, bold=True)


def rl_color(value):
    return colors.HexColor(f"#{value}")


def build_pdf():
    """Create a share-ready PDF companion and keep its layout independent of Word."""
    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.78 * inch,
        bottomMargin=0.72 * inch,
        title="Coter Pro - Dossier para Clinicas",
        author="Coter Pro",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="KickerCoter", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8.8,
        leading=11, textColor=rl_color(PURPLE), spaceAfter=8, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="CoverTitleCoter", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=28,
        leading=32, textColor=rl_color(NAVY), spaceAfter=9, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="CoverSubCoter", parent=styles["Normal"], fontName="Helvetica", fontSize=13.2,
        leading=18, textColor=rl_color(MUTED), spaceAfter=15, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="H1Coter", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16,
        leading=20, textColor=rl_color(NAVY), spaceBefore=0, spaceAfter=8, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="LeadCoter", parent=styles["Normal"], fontName="Helvetica", fontSize=12.3,
        leading=16, textColor=rl_color(NAVY), spaceAfter=14, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="FeatureTitleCoter", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12.1,
        leading=15, textColor=rl_color(PURPLE), spaceBefore=5, spaceAfter=2, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="FeatureBodyCoter", parent=styles["Normal"], fontName="Helvetica", fontSize=10.3,
        leading=13.8, textColor=rl_color(INK), spaceAfter=7, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="CalloutTitleCoter", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.9,
        leading=14, textColor=rl_color(NAVY), spaceAfter=2, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="CalloutBodyCoter", parent=styles["Normal"], fontName="Helvetica", fontSize=9.8,
        leading=13, textColor=rl_color(INK), alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="TableHeaderCoter", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8.8,
        leading=11, textColor=colors.white, alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="TableBodyCoter", parent=styles["Normal"], fontName="Helvetica", fontSize=8.75,
        leading=11.6, textColor=rl_color(INK), alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="TableLabelCoter", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9.1,
        leading=11.6, textColor=rl_color(NAVY), alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="SmallCoter", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=8.9,
        leading=11, textColor=rl_color(MUTED), alignment=TA_LEFT,
    ))
    styles.add(ParagraphStyle(
        name="EndCoter", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.2,
        leading=13, textColor=rl_color(NAVY), alignment=TA_CENTER,
    ))

    def p(text, style):
        return Paragraph(escape(text), styles[style])

    def p_rich(text, style):
        return Paragraph(text, styles[style])

    def feature(title, body, accent=PURPLE):
        title_html = f'<font color="#{accent}">{escape(title)}</font>'
        return KeepTogether([
            p_rich(title_html, "FeatureTitleCoter"),
            p(body, "FeatureBodyCoter"),
        ])

    def callout(title, body, fill=PALE_PURPLE):
        content = [p(title, "CalloutTitleCoter"), p(body, "CalloutBodyCoter")]
        table = RLTable([[content]], colWidths=[7.0 * inch])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), rl_color(fill)),
            ("BOX", (0, 0), (-1, -1), 0.6, rl_color(fill)),
            ("LEFTPADDING", (0, 0), (-1, -1), 14),
            ("RIGHTPADDING", (0, 0), (-1, -1), 14),
            ("TOPPADDING", (0, 0), (-1, -1), 11),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return table

    def data_table(headers, rows, widths):
        content = [[p(header, "TableHeaderCoter") for header in headers]]
        for row in rows:
            cells = []
            for item in row:
                cells.append(p(item, "TableBodyCoter"))
            content.append(cells)
        converted = [w / 1440 * inch for w in widths]
        table = RLTable(content, colWidths=converted, repeatRows=1, hAlign="LEFT")
        commands = [
            ("BACKGROUND", (0, 0), (-1, 0), rl_color(NAVY)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, rl_color(LIGHT_LINE)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, 0), 8),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
            ("TOPPADDING", (0, 1), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 8),
        ]
        for idx in range(1, len(content)):
            if idx % 2 == 0:
                commands.append(("BACKGROUND", (0, idx), (-1, idx), rl_color(PALE_BLUE)))
            else:
                commands.append(("BACKGROUND", (0, idx), (-1, idx), colors.white))
        table.setStyle(TableStyle(commands))
        return table

    def page_chrome(canvas, document):
        canvas.saveState()
        width, height = letter
        canvas.setStrokeColor(rl_color(LIGHT_LINE))
        canvas.setLineWidth(0.45)
        canvas.line(pdf.leftMargin, height - 0.48 * inch, width - pdf.rightMargin, height - 0.48 * inch)
        canvas.setFillColor(rl_color(MUTED))
        canvas.setFont("Helvetica-Bold", 7.8)
        canvas.drawString(pdf.leftMargin, height - 0.37 * inch, "COTER PRO / DOSSIER PARA CLINICAS")
        canvas.setFont("Helvetica", 8)
        canvas.drawString(pdf.leftMargin, 0.39 * inch, "Coter Pro - Continuidad terapeutica entre sesiones")
        canvas.drawRightString(width - pdf.rightMargin, 0.39 * inch, f"Pagina {document.page}")
        canvas.restoreState()

    story = []
    # Page 1 - cover
    story += [
        Spacer(1, 0.17 * inch),
        p("DOSSIER DE PRODUCTO PARA CLINICAS", "KickerCoter"),
        p("Continuidad terapeutica entre sesiones", "CoverTitleCoter"),
        p(
            "Una plataforma para que el equipo clinico acompañe el proceso entre consultas con mas estructura, contexto y claridad.",
            "CoverSubCoter",
        ),
        RLImage(str(BRAND_IMAGE), width=7.0 * inch, height=3.675 * inch),
        Spacer(1, 0.16 * inch),
        callout(
            "Una capa de continuidad para la practica clinica.",
            "Coter Pro centraliza check-ins, tareas terapeuticas, escalas, notas y comunicacion en un entorno pensado para que cada sesion empiece con la historia reciente del paciente.",
            PALE_CYAN,
        ),
        Spacer(1, 0.16 * inch),
        p("Preparado para conversaciones de demostracion y piloto - Julio de 2026", "SmallCoter"),
        PageBreak(),
    ]
    # Page 2 - problem
    story += [
        p("POR QUE EXISTE COTER PRO", "KickerCoter"),
        p("Lo importante no ocurre solo durante la sesion.", "H1Coter"),
        p(
            "Entre consultas aparecen avances, bloqueos, dudas y momentos relevantes. Sin un espacio comun, esa informacion llega tarde, se reparte en distintos canales o se pierde.",
            "LeadCoter",
        ),
        feature(
            "Una historia reciente, no solo el recuerdo de la ultima cita.",
            "Los check-ins breves ayudan a recoger estado de animo, ansiedad, energia y reflexiones cuando son relevantes; el profesional los revisa dentro del contexto del proceso.",
        ),
        feature(
            "Tareas que dejan de depender de mensajes y memoria.",
            "El equipo puede proponer una tarea concreta, con una estructura clara para el paciente y una respuesta que se recupera antes de la siguiente consulta.",
            CYAN,
        ),
        feature(
            "Un canal profesional separado de la mensajeria personal.",
            "Coter concentra la comunicacion vinculada al seguimiento en el mismo lugar que los objetivos, las notas y la actividad reciente.",
        ),
        Spacer(1, 0.12 * inch),
        callout(
            "La conversacion clinica sigue siendo el centro.",
            "Coter no sustituye la relacion terapeutica ni el juicio profesional: da una estructura ligera para que el tiempo entre sesiones aporte contexto util.",
        ),
        PageBreak(),
    ]
    # Page 3 - flow
    story += [
        p("COMO SE UTILIZA", "KickerCoter"),
        p("Un flujo directo para profesionales y pacientes.", "H1Coter"),
        p(
            "La experiencia esta diseñada para reducir friccion: el profesional define el seguimiento y el paciente accede con una interfaz sencilla.",
            "LeadCoter",
        ),
        data_table(
            ["Momento", "Que ocurre", "Resultado para la consulta"],
            [
                ["Preparacion", "El profesional crea el acceso, define objetivos y asigna tareas o cuestionarios acordes al caso.", "Un plan entre sesiones claro y personalizado."],
                ["Entre sesiones", "El paciente completa check-ins, ejercicios estructurados y escalas desde su espacio personal.", "Informacion ordenada, no mensajes dispersos."],
                ["Antes de la cita", "El profesional revisa actividad, respuestas, progreso y alertas de seguimiento desde el panel.", "La sesion puede empezar con contexto reciente."],
            ],
            [1680, 4260, 3420],
        ),
        Spacer(1, 0.14 * inch),
        feature(
            "Pensado para una adopcion gradual.",
            "Una clinica puede empezar con un grupo acotado de profesionales y pacientes, acordar que informacion es util y ajustar el flujo antes de ampliarlo.",
        ),
        feature(
            "Una experiencia paciente que no abruma.",
            "La persona encuentra una tarea visible, un check-in breve y el historial de su propio progreso. El objetivo es facilitar el uso real, no llenar su dia de pantallas.",
            CYAN,
        ),
        PageBreak(),
    ]
    # Page 4 - capabilities
    story += [
        p("CAPACIDADES DE PRODUCTO", "KickerCoter"),
        p("Las herramientas que sostienen el seguimiento.", "H1Coter"),
        p(
            "Coter Pro combina funciones practicas que el equipo puede utilizar segun su modelo terapeutico y sus protocolos internos.",
            "LeadCoter",
        ),
        feature("Check-ins emocionales", "Registro breve de animo, ansiedad, energia y pensamientos recientes para observar el intervalo entre consultas y enriquecer la conversacion clinica."),
        feature("Tareas terapeuticas estructuradas", "Biblioteca y tareas personalizadas para trabajo cognitivo-conductual, registros de pensamiento, activacion conductual, exposicion gradual y ejercicios clasicos.", CYAN),
        feature("Escalas y evolucion", "Cuestionarios como PHQ-9, GAD-7 y BDI-II con puntuacion e historial temporal, utilizados como apoyo a la valoracion profesional."),
        feature("Panel clinico con contexto", "Notas, objetivos, mensajes, actividad reciente y alertas de seguimiento reunidos en una unica vista para preparar el encuentro terapeutico.", CYAN),
        Spacer(1, 0.1 * inch),
        callout(
            "Informacion para revisar, no decisiones automaticas.",
            "Las escalas, tendencias y alertas son apoyos de seguimiento. La interpretacion, la priorizacion y cualquier decision clinica pertenecen siempre al profesional responsable.",
            PALE_CYAN,
        ),
        PageBreak(),
    ]
    # Page 5 - pilot
    story += [
        p("PILOTO CON UNA CLINICA", "KickerCoter"),
        p("Un despliegue pequeño, medible y adaptable.", "H1Coter"),
        p(
            "La mejor manera de valorar Coter es probarlo con un reto concreto: por ejemplo, mejorar la adherencia a tareas o llegar a cada sesion con una vision mas actualizada.",
            "LeadCoter",
        ),
        data_table(
            ["Etapa", "Trabajo conjunto", "Evidencia que se revisa"],
            [
                ["Definicion", "Seleccionar perfiles, protocolo de uso, responsables y limites del piloto.", "Objetivo operativo y criterios de exito acordados."],
                ["Configuracion", "Preparar tareas, escalas, mensajes de bienvenida y una orientacion breve al equipo.", "Flujo listo para un grupo inicial de pacientes."],
                ["Seguimiento", "Revisar uso, fricciones y aprendizaje con una cadencia acordada.", "Datos de adopcion y feedback cualitativo del equipo."],
                ["Decision", "Valorar ajustes, continuidad o ampliacion del alcance.", "Conclusiones utiles para la clinica y para Coter."],
            ],
            [1560, 4680, 3120],
        ),
        Spacer(1, 0.12 * inch),
        feature("Indicadores que una clinica puede acordar desde el inicio.", "Uso sostenido por pacientes, porcentaje de tareas completadas, frecuencia de check-ins, tiempo de preparacion de sesiones y valoracion del equipo sobre la calidad del contexto disponible."),
        feature("El piloto no exige cambiar todo de golpe.", "El alcance, el numero de profesionales, los perfiles de pacientes y la frecuencia de revision se definen de forma conjunta. Primero se valida la utilidad; despues se decide si tiene sentido escalar.", CYAN),
        PageBreak(),
    ]
    # Page 6 - trust and CTA
    story += [
        p("CONFIANZA Y SIGUIENTE PASO", "KickerCoter"),
        p("Tecnologia al servicio de una practica responsable.", "H1Coter"),
        p(
            "Coter se ha construido para trabajar con informacion sensible y para preservar la responsabilidad clinica en manos de las personas.",
            "LeadCoter",
        ),
        feature("Proteccion de datos sensibles", "El producto incorpora cifrado de datos sensibles en base de datos, sesiones seguras, controles de acceso separados entre profesional y paciente y registro de actividad relevante."),
        feature("Uso responsable", "Coter no diagnostica, no recomienda tratamientos y no sustituye los protocolos de crisis de la clinica. No debe utilizarse como canal de atencion urgente; cada despliegue debe definir sus circuitos asistenciales y de proteccion de datos.", CYAN),
        feature("Integracion con criterio clinico", "Las tareas, escalas, mensajes y alertas se configuran alrededor del enfoque y los limites de cada equipo. La herramienta acompaña el proceso; no impone un modelo terapeutico unico."),
        Spacer(1, 0.12 * inch),
        callout(
            "El siguiente paso: una conversacion de descubrimiento.",
            "Identifiquemos un reto de seguimiento que merezca la pena resolver, el grupo inicial de pacientes y las condiciones necesarias para un piloto seguro y util. Solicita una demostracion en coter.app.",
        ),
        Spacer(1, 0.2 * inch),
        p("COTER PRO - CONTINUIDAD TERAPEUTICA ENTRE SESIONES", "EndCoter"),
    ]
    pdf.build(story, onFirstPage=page_chrome, onLaterPages=page_chrome)
    print(PDF_PATH)


def main():
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Coter Pro - Dossier para Clinicas"
    props.subject = "Dossier informativo de Coter Pro para clinicas"
    props.author = "Coter Pro"
    props.company = "Coter Pro"
    props.comments = "Dossier comercial de producto."

    add_cover(doc)
    page_break(doc)
    add_problem_page(doc)
    page_break(doc)
    add_workflow_page(doc)
    page_break(doc)
    add_capabilities_page(doc)
    page_break(doc)
    add_pilot_page(doc)
    page_break(doc)
    add_trust_page(doc)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    build_pdf()


if __name__ == "__main__":
    main()
